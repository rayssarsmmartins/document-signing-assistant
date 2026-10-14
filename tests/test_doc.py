from cs import doc


def test_pdf_with_text_is_readable(contract_pdf):
    info, text = doc.inspect(contract_pdf)
    assert info.status == "ok" and info.can_read and info.can_send
    assert info.next == "read"
    assert "maria.souza@example.com" in text
    assert info.pages == 1


def test_scanned_pdf_goes_to_ask_signers(blank_pdf):
    info, _ = doc.inspect(blank_pdf)
    assert info.status == "no_text"
    assert info.can_send and not info.can_read
    assert info.next == "ask_signers"


def test_encrypted_pdf_asks_new_file(encrypted_pdf):
    info, text = doc.inspect(encrypted_pdf)
    assert info.status == "encrypted"
    assert not info.can_send
    assert info.next == "ask_new_file"
    assert text == ""


def test_unsupported_extension(tmp_path):
    f = tmp_path / "planilha.xlsx"
    f.write_bytes(b"x")
    info, _ = doc.inspect(f)
    assert info.status == "unsupported_format" and info.next == "ask_new_file"


def test_corrupt_pdf(tmp_path):
    f = tmp_path / "quebrado.pdf"
    f.write_bytes(b"isto nao e um pdf")
    info, _ = doc.inspect(f)
    assert info.status == "corrupt" and info.next == "ask_new_file"


def test_docx_text(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("LOCADOR: Carlos Alberto Nunes, carlos@example.com")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "LOCATARIA"
    t.rows[0].cells[1].text = "Ana Beatriz Rocha"
    p = tmp_path / "locacao.docx"
    d.save(str(p))
    info, text = doc.inspect(p)
    assert info.status == "ok"
    assert "Ana Beatriz Rocha" in text and "carlos@example.com" in text


def test_image_is_sendable_but_not_readable(tmp_path):
    f = tmp_path / "foto.jpg"
    f.write_bytes(b"\xff\xd8\xff")
    info, _ = doc.inspect(f)
    assert info.can_send and not info.can_read and info.next == "ask_signers"


def test_legacy_doc(tmp_path):
    f = tmp_path / "antigo.doc"
    f.write_bytes(b"\xd0\xcf\x11\xe0")
    info, _ = doc.inspect(f)
    assert info.status == "unreadable_format" and info.next == "ask_signers"
